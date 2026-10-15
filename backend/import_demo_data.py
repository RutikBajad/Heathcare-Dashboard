import argparse
from datetime import date, datetime, time, timezone
import os
from pathlib import Path
from urllib.parse import urlparse

from dotenv import load_dotenv
from openpyxl import load_workbook
from pymongo import MongoClient
from pymongo.errors import PyMongoError
from werkzeug.security import generate_password_hash


ROOT_DIR = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = ROOT_DIR / 'healthcare_demo_data_aug_to_nov_2026.xlsx'
EXPECTED_COUNTS = {
    'doctors': 12,
    'staff_accounts': 20,
    'diseases': 25,
    'patients': 160,
    'patient_logins': 24,
    'appointments': 170,
}
STAFF_PAGES = ['overview', 'patients', 'doctors', 'appointments', 'add-record', 'analytics']
PATIENT_PAGES = ['appointments', 'hospital-info', 'doctors', 'settings']


def text(value, field):
    if value is None or not str(value).strip():
        raise ValueError(f'{field} is required in the demo workbook.')
    return str(value).strip()


def whole_number(value, field):
    try:
        number = int(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f'{field} must be a whole number.') from error
    if isinstance(value, str) and str(number) != value.strip():
        raise ValueError(f'{field} must be a whole number.')
    if isinstance(value, (int, float)) and number != value:
        raise ValueError(f'{field} must be a whole number.')
    return number


def iso_date(value, field):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    try:
        return date.fromisoformat(text(value, field)).isoformat()
    except ValueError as error:
        raise ValueError(f'{field} must be a valid date.') from error


def iso_time(value, field):
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.strftime('%H:%M')
    try:
        return time.fromisoformat(text(value, field)).strftime('%H:%M')
    except ValueError as error:
        raise ValueError(f'{field} must be a valid time.') from error


def read_rows(workbook, sheet_name, expected_headers):
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f'The demo workbook is missing its {sheet_name} sheet.')
    rows = workbook[sheet_name].iter_rows(values_only=True)
    headers = next(rows, None)
    if tuple(headers or ()) != tuple(expected_headers):
        raise ValueError(f'The {sheet_name} sheet has unexpected columns.')
    return [dict(zip(expected_headers, row)) for row in rows if any(value is not None for value in row)]


def load_demo_records(workbook_path):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        intro = workbook['Start Here'].cell(row=2, column=1).value
        if not isinstance(intro, str) or not intro.startswith('Demo-only fictional data.'):
            raise ValueError('Refusing to import a workbook not marked as fictional demo data.')

        doctors_rows = read_rows(workbook, 'Doctors', (
            'Reference ID', 'Name', 'Specialty', 'Department', 'City', 'Availability',
        ))
        staff_rows = read_rows(workbook, 'Staff Accounts', (
            'Reference ID', 'Name', 'Email / Account ID', 'Temporary Password', 'Role',
        ))
        disease_rows = read_rows(workbook, 'Diseases', (
            'Disease ID', 'Name', 'Category', 'Description',
        ))
        patient_rows = read_rows(workbook, 'Patients', (
            'Reference ID', 'Name', 'Age', 'City', 'Gender', 'Blood Group', 'Status',
            'Last Visit', 'Risk', 'Conditions (Disease IDs)', 'Portal Login',
        ))
        login_rows = read_rows(workbook, 'Patient Logins', (
            'Patient Reference ID', 'Name', 'Age', 'City', 'Gender', 'Demo Email',
            'Demo Password', 'Role',
        ))
        appointment_rows = read_rows(workbook, 'Appointments', (
            'Reference ID', 'Patient Reference ID', 'Patient', 'Patient Demo Email (online only)',
            'Doctor', 'Date', 'Time', 'Status', 'Reason', 'Booking Channel', 'Entered By',
        ))
    finally:
        workbook.close()

    actual_counts = {
        'doctors': len(doctors_rows),
        'staff_accounts': len(staff_rows),
        'diseases': len(disease_rows),
        'patients': len(patient_rows),
        'patient_logins': len(login_rows),
        'appointments': len(appointment_rows),
    }
    if actual_counts != EXPECTED_COUNTS:
        raise ValueError(f'Demo workbook counts do not match the expected dataset: {actual_counts}.')
    for label, rows, id_field in (
        ('doctor', doctors_rows, 'Reference ID'),
        ('staff account', staff_rows, 'Reference ID'),
        ('disease', disease_rows, 'Disease ID'),
        ('patient login', login_rows, 'Patient Reference ID'),
        ('appointment', appointment_rows, 'Reference ID'),
    ):
        identifiers = [text(row[id_field], f'{label.title()} ID') for row in rows]
        if len(identifiers) != len(set(identifiers)):
            raise ValueError(f'{label.title()} IDs must be unique.')

    doctors = [{
        'id': text(row['Reference ID'], 'Doctor reference ID'),
        'name': text(row['Name'], 'Doctor name'),
        'specialty': text(row['Specialty'], 'Doctor specialty'),
        'department': text(row['Department'], 'Doctor department'),
        'city': text(row['City'], 'Doctor city'),
        'availability': text(row['Availability'], 'Doctor availability'),
        'active': True,
        'demoRecord': True,
    } for row in doctors_rows]

    diseases = [{
        'id': text(row['Disease ID'], 'Disease ID'),
        'name': text(row['Name'], 'Disease name'),
        'category': text(row['Category'], 'Disease category'),
        'description': text(row['Description'], 'Disease description'),
        'demoRecord': True,
    } for row in disease_rows]

    patient_ids = {text(row['Reference ID'], 'Patient reference ID') for row in patient_rows}
    disease_ids = {disease['id'] for disease in diseases}
    doctor_names = {doctor['name'] for doctor in doctors}

    patients = []
    for row in patient_rows:
        conditions = [
            value.strip()
            for value in str(row['Conditions (Disease IDs)'] or '').split(',')
            if value.strip()
        ]
        if not set(conditions) <= disease_ids:
            raise ValueError('A patient refers to a disease missing from the Diseases sheet.')
        age = whole_number(row['Age'], 'Patient age')
        if not 0 <= age <= 120:
            raise ValueError('A patient age is outside the supported range.')
        patients.append({
            'id': text(row['Reference ID'], 'Patient reference ID'),
            'name': text(row['Name'], 'Patient name'),
            'age': age,
            'city': text(row['City'], 'Patient city'),
            'gender': text(row['Gender'], 'Patient gender'),
            'bloodGroup': text(row['Blood Group'], 'Patient blood group'),
            'status': text(row['Status'], 'Patient status'),
            'lastVisit': iso_date(row['Last Visit'], 'Patient last visit'),
            'risk': text(row['Risk'], 'Patient risk'),
            'conditions': conditions,
            'active': True,
            'portalLogin': str(row['Portal Login']).strip().lower() == 'yes',
            'demoRecord': True,
        })
    if len(patient_ids) != len(patients):
        raise ValueError('Patient reference IDs must be unique.')

    accounts = []
    users = []
    seen_emails = set()
    for row in staff_rows:
        email = text(row['Email / Account ID'], 'Staff account ID').lower()
        if email.rsplit('@', 1)[-1] != 'example.com' or email in seen_emails:
            raise ValueError('Staff accounts must have unique reserved example.com addresses.')
        seen_emails.add(email)
        if text(row['Role'], 'Staff role').lower() != 'staff':
            raise ValueError('Only staff accounts are supported in the Staff Accounts sheet.')
        ref = text(row['Reference ID'], 'Staff reference ID')
        name = text(row['Name'], 'Staff name')
        password = text(row['Temporary Password'], 'Staff temporary password')
        if len(password) < 8:
            raise ValueError('A demo staff password is shorter than the minimum supported length.')
        account_id = f'USR-{ref}'
        accounts.append({
            'id': account_id,
            'name': name,
            'email': email,
            'passwordHash': generate_password_hash(password, method='scrypt'),
            'role': 'staff',
            'active': True,
            'doctorName': None,
            'allowedPages': STAFF_PAGES,
            'createdAt': datetime.now(timezone.utc).isoformat(),
            'demoRecord': True,
        })
        users.append({
            'id': account_id,
            'accountId': account_id,
            'name': name,
            'email': email,
            'role': 'Staff',
            'active': True,
            'demoRecord': True,
        })

    for row in login_rows:
        patient_id = text(row['Patient Reference ID'], 'Patient login reference ID')
        if patient_id not in patient_ids:
            raise ValueError('A patient login refers to a patient missing from the Patients sheet.')
        email = text(row['Demo Email'], 'Patient demo email').lower()
        if email.rsplit('@', 1)[-1] != 'example.com' or email in seen_emails:
            raise ValueError('Demo accounts must have unique reserved example.com addresses.')
        seen_emails.add(email)
        if text(row['Role'], 'Patient role').lower() != 'patient':
            raise ValueError('Only patient accounts are supported in the Patient Logins sheet.')
        patient = next(item for item in patients if item['id'] == patient_id)
        if text(row['Name'], 'Patient login name') != patient['name']:
            raise ValueError('A patient login name does not match its linked patient record.')
        if whole_number(row['Age'], 'Patient login age') != patient['age']:
            raise ValueError('A patient login age does not match its linked patient record.')
        password = text(row['Demo Password'], 'Patient demo password')
        if len(password) < 8:
            raise ValueError('A demo patient password is shorter than the minimum supported length.')
        accounts.append({
            'id': f'USR-{patient_id}',
            'name': patient['name'],
            'email': email,
            'passwordHash': generate_password_hash(password, method='scrypt'),
            'role': 'patient',
            'active': True,
            'patientId': patient_id,
            'allowedPages': PATIENT_PAGES,
            'createdAt': datetime.now(timezone.utc).isoformat(),
            'demoRecord': True,
        })

    appointments = []
    for row in appointment_rows:
        patient_id = text(row['Patient Reference ID'], 'Appointment patient reference ID')
        doctor_name = text(row['Doctor'], 'Appointment doctor')
        if patient_id not in patient_ids or doctor_name not in doctor_names:
            raise ValueError('An appointment refers to a missing patient or doctor.')
        appointments.append({
            'id': text(row['Reference ID'], 'Appointment reference ID'),
            'patientId': patient_id,
            'patient': text(row['Patient'], 'Appointment patient name'),
            'doctor': doctor_name,
            'date': iso_date(row['Date'], 'Appointment date'),
            'time': iso_time(row['Time'], 'Appointment time'),
            'status': text(row['Status'], 'Appointment status'),
            'reason': text(row['Reason'], 'Appointment reason'),
            'bookingChannel': text(row['Booking Channel'], 'Appointment booking channel'),
            'enteredBy': text(row['Entered By'], 'Appointment staff name'),
            'demoRecord': True,
        })

    return {
        'diseases': diseases,
        'doctors': doctors,
        'patients': patients,
        'appointments': appointments,
        'accounts': accounts,
        'users': users,
    }


def validate_target(uri, database_name):
    parsed = urlparse(uri)
    host = (parsed.hostname or '').lower()
    if parsed.scheme not in {'mongodb', 'mongodb+srv'} or not host:
        raise ValueError('Set a valid hosted MongoDB URI in the ignored local .env file.')
    if host in {'localhost', '127.0.0.1', '::1'}:
        raise ValueError('Refusing to import into localhost; set the hosted MongoDB URI first.')
    if database_name != 'healthpilot':
        raise ValueError('Refusing to import: the target database must be healthpilot.')


def preflight_database(database, records_by_collection):
    for collection_name, documents in records_by_collection.items():
        collection = database[collection_name]
        for document in documents:
            existing = collection.find_one({'id': document['id']}, {'_id': 0, 'id': 1, 'demoRecord': 1, 'passwordHash': 1})
            if existing and existing.get('demoRecord') is not True:
                raise ValueError(f'Existing non-demo record conflicts with an import ID in {collection_name}.')
            if collection_name == 'accounts':
                existing_email = collection.find_one(
                    {'email': document['email']},
                    {'_id': 0, 'id': 1, 'demoRecord': 1},
                )
                if existing_email and (
                    existing_email.get('id') != document['id']
                    or existing_email.get('demoRecord') is not True
                ):
                    raise ValueError('An account email conflicts with an existing non-demo or different account.')


def import_records(database, records_by_collection):
    database.accounts.create_index('email', unique=True)
    preflight_database(database, records_by_collection)
    results = {}
    for collection_name, documents in records_by_collection.items():
        collection = database[collection_name]
        inserted = 0
        updated = 0
        for document in documents:
            existing = collection.find_one({'id': document['id']}, {'_id': 0, 'passwordHash': 1})
            if collection_name == 'accounts' and existing and existing.get('passwordHash'):
                document = {**document, 'passwordHash': existing['passwordHash']}
            result = collection.update_one({'id': document['id']}, {'$set': document}, upsert=True)
            inserted += int(result.upserted_id is not None)
            updated += int(result.upserted_id is None)
        results[collection_name] = {'inserted': inserted, 'updated': updated}
    return results


def main():
    parser = argparse.ArgumentParser(description='Import the fictional TakeCare demo workbook into hosted MongoDB.')
    parser.add_argument('--workbook', type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument('--apply', action='store_true', help='Write records to the hosted database; default is dry-run.')
    args = parser.parse_args()

    load_dotenv(ROOT_DIR / '.env')
    uri = os.getenv('MONGODB_URI', '')
    database_name = os.getenv('MONGODB_DATABASE', 'healthpilot')
    validate_target(uri, database_name)
    records_by_collection = load_demo_records(args.workbook)
    counts = {name: len(documents) for name, documents in records_by_collection.items()}

    if not args.apply:
        print(f'Dry run only; no database connection or writes. Target database: {database_name}.')
        print(f'Validated demo record counts: {counts}')
        print('Run again with --apply to write these records.')
        return

    client = MongoClient(uri, serverSelectionTimeoutMS=5000)
    try:
        client.admin.command('ping')
        results = import_records(client[database_name], records_by_collection)
        print(f'Demo import completed in database {database_name}: {results}')
    finally:
        client.close()


if __name__ == '__main__':
    try:
        main()
    except (OSError, ValueError, PyMongoError) as error:
        raise SystemExit(f'Demo import failed: {error}') from error
